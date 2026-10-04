"""Read plain text out of an uploaded CV. Checks the file's real type, not just its name."""

import io
import zipfile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import AppError

MAX_CV_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_TEXT_CHARS = 100_000

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class CvFileError(AppError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(422, code, message)


def detect_type(data: bytes) -> str:
    """PDF or DOCX by magic bytes. Anything else is refused."""
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return DOCX
        except zipfile.BadZipFile:
            pass
    raise CvFileError("unsupported_file", "Upload your CV as a PDF or a Word (.docx) file.")


def extract_text(data: bytes, content_type: str) -> str:
    try:
        text = _pdf_text(data) if content_type == PDF else _docx_text(data)
    except CvFileError:
        raise
    except Exception as exc:  # A damaged file can fail in many library-specific ways.
        raise CvFileError(
            "unreadable_file", "We couldn't read this file. Try exporting it again as a PDF."
        ) from exc
    text = text.strip()
    if not text:
        raise CvFileError(
            "no_text",
            "This file has no readable text. It may be a scan; try a PDF exported from a "
            "word processor.",
        )
    return text[:MAX_TEXT_CHARS]


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise CvFileError("unreadable_file", "We couldn't open this PDF.") from exc
    if reader.is_encrypted:
        raise CvFileError("encrypted_file", "This PDF is password-protected. Remove the password.")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise CvFileError("too_many_pages", f"CVs are limited to {MAX_PDF_PAGES} pages.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
